from datetime import date
from io import BytesIO
import unittest
import pandas as pd
from openpyxl import load_workbook as open_xlsx
from reda_queue import load_workbook, new_notes, queue_view, export_workbook, parse_request_date


def workbook(frame):
    output = BytesIO()
    frame.to_excel(output, index=False)
    return output.getvalue()


class QueueTests(unittest.TestCase):
    def test_missing_fields_are_reported_not_invented(self):
        source, notes, mapping, missing = load_workbook(workbook(pd.DataFrame({'Outro': ['X', None]})))
        self.assertEqual(len(missing), 5)
        self.assertEqual(mapping, {})
        self.assertEqual(list(source.columns), ['Outro'])
        self.assertTrue(notes['ESTADO LOCAL'].eq('').all())
        self.assertTrue(notes['PRIORIDADE LOCAL'].eq('').all())
        self.assertEqual(notes['ID LOCAL'].nunique(), len(source))

    def test_normalized_headers_and_incomplete_values(self):
        source, notes, mapping, missing = load_workbook(workbook(pd.DataFrame({' titulo ': ['A', None], 'Autor/Fonte': [None, 'B'], 'Outra': [1, 2]})))
        self.assertEqual(mapping['TÍTULO'], ' titulo ')
        view = queue_view(source, notes, mapping)
        self.assertIn('AUTOR/FONTE', view.loc[0, 'CAMPOS POR COMPLETAR / CONFIRMAR'])
        self.assertIn('Sem título', view.loc[1, 'TÍTULO / REFERÊNCIA'])
        self.assertEqual(source.loc[0, 'Outra'], 1)

    def test_priority_then_age_then_original_order(self):
        source = pd.DataFrame({'Data': ['01/01/2026', '01/10/2026', None, '01/01/2025', '01/10/2026']})
        notes = new_notes(5)
        notes['PRIORIDADE LOCAL'] = ['Normal', 'Urgente', 'Urgente', '', 'Urgente']
        result = queue_view(source, notes, {}, 'Data', today=date(2026, 10, 26))
        self.assertEqual(list(result.index), [1, 4, 2, 0, 3])
        self.assertEqual(result.loc[0, 'PEDIDO ANTIGO'], 'Sim')
        self.assertEqual(result.loc[2, 'PEDIDO ANTIGO'], 'Por confirmar')

    def test_dates_unknown_invalid_future_and_numeric(self):
        source = pd.DataFrame({'Data': ['31/02/2026', 46000, '2027-01-01', '', '2026-10-25']})
        view = queue_view(source, new_notes(5), {}, 'Data', today=date(2026, 10, 26))
        self.assertTrue(view.loc[0:3, 'DIAS EM ESPERA'].isna().all())
        self.assertEqual(view.loc[4, 'DIAS EM ESPERA'], 1)
        self.assertEqual(parse_request_date(date(2026, 1, 1))[0], date(2026, 1, 1))

    def test_export_restore_preserves_source_notes_ids_and_extra_columns(self):
        source = pd.DataFrame({'TÍTULO': ['Mesmo', 'Mesmo'], 'Extra': ['a', 'b']})
        notes = new_notes(2)
        notes.loc[1, ['BLOQUEIOS', 'PRÓXIMO PASSO']] = ['Falta licença', 'Consultar autor']
        view = queue_view(source, notes, {})
        result = export_workbook(source, notes, view)
        restored, saved, _, _ = load_workbook(result, 'Recursos', True)
        pd.testing.assert_frame_equal(restored, source, check_dtype=False)
        pd.testing.assert_frame_equal(saved, notes, check_dtype=False)
        _, fresh, _, _ = load_workbook(result, 'Recursos', False)
        self.assertTrue(fresh['BLOQUEIOS'].eq('').all())

    def test_source_state_never_becomes_local_state(self):
        source = pd.DataFrame({'Estado': ['Aprovado', 'Pendente']})
        notes = new_notes(2)
        view = queue_view(source, notes, {}, state_column='Estado')
        self.assertEqual(list(view['ESTADO NA FONTE']), ['Aprovado', 'Pendente'])
        self.assertTrue(view['ESTADO LOCAL'].eq('').all())

    def test_invalid_workbook_and_invalid_saved_queue(self):
        with self.assertRaises(Exception):
            load_workbook(b'not excel')
        with self.assertRaises(ValueError):
            load_workbook(workbook(pd.DataFrame({'A': [1]})), 0, True)

    def test_empty_sheet(self):
        source, notes, mapping, _ = load_workbook(workbook(pd.DataFrame(columns=['TÍTULO'])))
        self.assertTrue(queue_view(source, notes, mapping).empty)

    def test_export_does_not_create_formulas(self):
        source = pd.DataFrame({'TÍTULO': ['=1+1']})
        notes = new_notes(1)
        raw = export_workbook(source, notes, queue_view(source, notes, {}))
        self.assertEqual(open_xlsx(BytesIO(raw))['Recursos']['A2'].data_type, 's')
